# Зависимости: pip install openpyxl cryptography
# Использование: python unveil.py <carrier.xlsx> <PRIVATE_KEY|@file|$ENV> <output.xlsx> [--mode production|demo] [--sheet NAME] [--start-cell A1] [--print]

import argparse
import base64
import hashlib
import hmac
import math
import os
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.utils.cell import range_boundaries, coordinate_from_string
    from openpyxl.utils import column_index_from_string
except ImportError:
    sys.exit("openpyxl не установлен: pip install openpyxl")

try:
    from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
    from cryptography.hazmat.primitives.asymmetric.x25519 import (
        X25519PrivateKey, X25519PublicKey,
    )
    from cryptography.hazmat.primitives.serialization import Encoding, PublicFormat
except ImportError:
    sys.exit("cryptography не установлен: pip install cryptography")

MAGIC          = b"CVX3"
VERSION        = 1
MODE_DEMO      = 1
MODE_PROD      = 2
FONTS          = ["Arial", "Calibri", "Segoe UI", "Tahoma"]
MAX_BYTES      = 1_048_576
MAX_CELLS      = 200_000
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLS = 16_384


def cv_hash(data: bytes) -> bytes:
    return hashlib.sha256(data).digest()


def cv_hmac(key: bytes, data: bytes) -> bytes:
    return hmac.new(key, data, hashlib.sha256).digest()


def cv_ctr(data: bytes, key: bytes, nonce: bytes) -> bytes:
    cipher = Cipher(algorithms.AES(key), modes.CTR(nonce))
    enc = cipher.encryptor()
    return enc.update(data) + enc.finalize()


def cv_kdf(secret_value: bytes, ep: bytes, recipient: bytes, mode: int) -> bytes:
    prk  = cv_hmac(b"\x00" * 32, secret_value)
    info = f"CellVeilXL-v3/{mode}".encode("ascii")
    return cv_hmac(prk, info + ep + recipient + b"\x01")


def cv_b64url_decode(value: str, size: int) -> bytes:
    value = value.replace(" ", "").replace("\r", "").replace("\n", "").strip()
    expected_len = (size * 8 + 5) // 6
    if len(value) != expected_len:
        raise ValueError(
            f"Длина ключа не соответствует режиму: ожидалось {expected_len}, "
            f"получено {len(value)}"
        )
    pad = (-len(value)) % 4
    decoded = base64.urlsafe_b64decode(value + "=" * pad)
    result  = decoded[:size]
    reencoded = base64.urlsafe_b64encode(result).rstrip(b"=").decode("ascii")
    if reencoded != value:
        raise ValueError("Неканоническое кодирование ключа (лишние биты)")
    return result


def cv_public_x25519(priv: bytes) -> bytes:
    k = X25519PrivateKey.from_private_bytes(priv)
    return k.public_key().public_bytes(Encoding.Raw, PublicFormat.Raw)


def cv_x25519(priv: bytes, point: bytes) -> bytes:
    k = X25519PrivateKey.from_private_bytes(priv)
    p = X25519PublicKey.from_public_bytes(point)
    return k.exchange(p)


DEMO_P = 154742504910672534362390399
DEMO_Q = 77371252455336267181195199


def demo_public(priv: bytes) -> bytes:
    n = int.from_bytes(priv, "little")
    return pow(4, n, DEMO_P).to_bytes(11, "little")


def demo_shared(priv: bytes, pub: bytes) -> bytes:
    n = int.from_bytes(priv, "little")
    y = int.from_bytes(pub, "little")
    if not (2 <= n < DEMO_Q):
        raise ValueError("Неверный demo закрытый ключ")
    if not (2 <= y < DEMO_P - 1):
        raise ValueError("Неверный demo открытый ключ")
    if pow(y, DEMO_Q, DEMO_P) != 1:
        raise ValueError("Demo-ключ вне подгруппы")
    return pow(y, n, DEMO_P).to_bytes(11, "little")


def cv_public(priv: bytes, mode: int) -> bytes:
    if mode == MODE_PROD:  return cv_public_x25519(priv)
    if mode == MODE_DEMO:  return demo_public(priv)
    raise ValueError(f"Неизвестный режим {mode}")


def shared_key(priv: bytes, pub: bytes, mode: int) -> bytes:
    if mode == MODE_PROD:   b = cv_x25519(priv, pub)
    elif mode == MODE_DEMO: b = demo_shared(priv, pub)
    else: raise ValueError(f"Неизвестный режим {mode}")
    if not any(b):
        raise ValueError("Нулевой общий секрет")
    return b


def cv_open(payload: bytes, private_text: str, mode: int) -> bytes:
    if   mode == MODE_PROD: size = 32
    elif mode == MODE_DEMO: size = 11
    else: raise ValueError(f"Неизвестный режим {mode}")

    priv      = cv_b64url_decode(private_text, size)
    recipient = cv_public(priv, mode)
    if len(payload) < size + 16:
        raise ValueError("Пакет короче заголовка (ep + nonce)")
    ep        = payload[:size]
    nonce     = payload[size:size + 16]
    encrypted = payload[size + 16:]

    secret    = shared_key(priv, ep, mode)
    key       = cv_kdf(secret, ep, recipient, mode)
    return cv_ctr(encrypted, key, nonce)


def rotation_xlsx_to_vba(rotation):
    if rotation is None or rotation == 0:
        return 0
    if 1 <= rotation <= 90:
        return rotation
    if 91 <= rotation <= 180:
        return -(rotation - 90)
    return 255


def orientation_base(vba_rot: int) -> int:
    if vba_rot == 0:
        return 0
    base = math.floor(vba_rot / 4) * 4
    return max(-88, min(84, base))


def _rgb_from_color(color):
    if color is None:
        return (0, 0, 0)
    rgb = getattr(color, "rgb", None)
    if rgb is None:
        return (0, 0, 0)
    if isinstance(rgb, str) and len(rgb) == 8:
        return (int(rgb[2:4], 16), int(rgb[4:6], 16), int(rgb[6:8], 16))
    if isinstance(rgb, int):
        return ((rgb >> 16) & 0xFF, (rgb >> 8) & 0xFF, rgb & 0xFF)
    return (0, 0, 0)


def _rgb_from_font(font):
    if font is None:
        return (0, 0, 0)
    return _rgb_from_color(font.color)


def _cell_text(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, CellRichText):
        return "".join(b if isinstance(b, str) else b.text for b in v)
    return str(v)


def _char_colors(cell):
    v = cell.value
    if isinstance(v, CellRichText):
        out = []
        for block in v:
            if isinstance(block, str):
                out.extend([_rgb_from_font(cell.font)] * len(block))
            elif isinstance(block, TextBlock):
                f = block.font if block.font else cell.font
                out.extend([_rgb_from_font(f)] * len(block.text))
            else:
                out.extend([(0, 0, 0)] * len(str(block)))
        return out
    return [_rgb_from_font(cell.font)] * len(str(v))


def _font_index(cell) -> int:
    name = cell.font.name if cell.font else None
    if not name:
        name = "Calibri"
    for i, f in enumerate(FONTS):
        if f.lower() == name.lower():
            return i
    raise ValueError(f"Неизвестный шрифт-носитель: {name!r}")


def _is_merged(ws, cell) -> bool:
    for rng in ws.merged_cells.ranges:
        if (rng.min_row <= cell.row <= rng.max_row
                and rng.min_col <= cell.column <= rng.max_col):
            return True
    return False


def _is_usable(ws, cell) -> bool:
    if _is_merged(ws, cell):
        return False
    if cell.data_type in ("f", "e"):
        return False
    if not isinstance(cell.value, (str, CellRichText)):
        return False
    if len(_cell_text(cell)) == 0:
        return False
    if cell.font and cell.font.bold:
        return False
    rot = rotation_xlsx_to_vba(
        cell.alignment.textRotation if cell.alignment else None)
    if rot != 0 and (rot < -90 or rot > 90):
        return False
    return True


def _iter_usable_cells(ws):
    min_col, min_row, max_col, max_row = range_boundaries(ws.calculate_dimension())
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(row=r, column=c)
            if _is_usable(ws, cell):
                yield cell


class BitReader:
    def __init__(self, ws):
        self._it = _iter_usable_cells(ws)
        self._bits = []
        self._pos = 0

    def _fill(self):
        while self._pos >= len(self._bits):
            try:
                cell = next(self._it)
            except StopIteration:
                raise ValueError("Носитель исчерпан (оборванный поток)")
            self._bits = self._cell_bits(cell)
            self._pos = 0

    def read_bits(self, n: int) -> int:
        v = 0
        for _ in range(n):
            self._fill()
            v = (v << 1) | self._bits[self._pos]
            self._pos += 1
        return v

    def read_bytes(self, n: int) -> bytes:
        return bytes(self.read_bits(8) for _ in range(n))

    @staticmethod
    def _cell_bits(cell):
        bits = []
        rot = rotation_xlsx_to_vba(
            cell.alignment.textRotation if cell.alignment else None)
        low2 = rot - orientation_base(rot)
        bits.append((low2 >> 1) & 1)
        bits.append(low2 & 1)
        fi = _font_index(cell)
        bits.append((fi >> 1) & 1)
        bits.append(fi & 1)
        for r, g, b in _char_colors(cell):
            low6 = ((r & 3) << 4) | ((g & 3) << 2) | (b & 3)
            for sh in (5, 4, 3, 2, 1, 0):
                bits.append((low6 >> sh) & 1)
        return bits


def read_packet(reader: BitReader) -> tuple[int, bytes]:
    header = reader.read_bytes(10)
    if header[:4] != MAGIC:
        raise ValueError(f"Неверная сигнатура: {header[:4]!r}, ожидалось {MAGIC!r}")
    mode, ver = header[4], header[5]
    if ver != VERSION:
        raise ValueError(f"Неподдерживаемая версия пакета: {ver}")
    length = int.from_bytes(header[6:10], "little")
    if not (28 <= length <= MAX_BYTES):
        raise ValueError(f"Некорректная длина пакета: {length}")
    return mode, reader.read_bytes(length)


def _read_field(s: str, pos: int) -> tuple[str, int]:
    colon = s.find(":", pos)
    if colon == -1 or colon - pos > 8:
        raise ValueError("Некорректная запись диапазона")
    token = s[pos:colon]
    if not token or not token.isdigit():
        raise ValueError("Некорректная длина поля")
    n = int(token)
    if n > len(s) - colon - 1:
        raise ValueError("Обрезанная запись диапазона")
    return s[colon + 1: colon + 1 + n], colon + 1 + n


def deserialize(data: bytes) -> tuple[int, int, list]:
    s = data.decode("utf-16-le")
    pos = 0
    h_str, pos = _read_field(s, pos)
    w_str, pos = _read_field(s, pos)
    if not h_str.isdigit() or not w_str.isdigit():
        raise ValueError("Некорректные размеры диапазона")
    h, w = int(h_str), int(w_str)
    if not (1 <= h <= EXCEL_MAX_ROWS and 1 <= w <= EXCEL_MAX_COLS):
        raise ValueError(f"Недопустимые размеры: {h}x{w}")
    if h * w > MAX_CELLS:
        raise ValueError(f"Диапазон слишком велик: {h}x{w}")

    cells = []
    for _ in range(h * w):
        kind, pos = _read_field(s, pos)
        val,  pos = _read_field(s, pos)
        fmt,  pos = _read_field(s, pos)
        if kind == "z":
            if val != "": raise ValueError("Некорректная пустая ячейка")
        elif kind == "s":
            pass
        elif kind == "b":
            if val not in ("0", "1"): raise ValueError("Некорректный Boolean")
        elif kind == "n":
            try: float(val)
            except ValueError: raise ValueError(f"Некорректное число: {val!r}")
        else:
            raise ValueError(f"Неизвестный тип ячейки: {kind!r} "
                             "(неверный ключ или повреждённые данные)")
        cells.append((kind, val, fmt))
    if pos != len(s):
        raise ValueError("Лишние данные после диапазона")
    return h, w, cells


def write_output(h: int, w: int, cells, path: str, start_cell: str = "A1"):
    col_letter, row = coordinate_from_string(start_cell)
    start_col = column_index_from_string(col_letter)
    if row + h - 1 > EXCEL_MAX_ROWS or start_col + w - 1 > EXCEL_MAX_COLS:
        raise ValueError("Выходной диапазон выходит за границы листа")

    wb = Workbook()
    ws = wb.active
    ws.title = "Restored"

    for i, (kind, val, fmt) in enumerate(cells):
        r = row + i // w
        c = start_col + i % w
        cell = ws.cell(row=r, column=c)
        if kind == "z":
            pass
        elif kind == "s":
            cell.value = val
            cell.number_format = "@" if not fmt else fmt
        elif kind == "b":
            cell.value = (val == "1")
        elif kind == "n":
            cell.value = float(val)
        cell.number_format = fmt

    wb.save(path)


def print_cells(h: int, w: int, cells):
    from openpyxl.utils import get_column_letter
    for i, (k, v, f) in enumerate(cells):
        addr = f"{get_column_letter(i % w + 1)}{i // w + 1}"
        print(f"  {addr}: kind={k!r:>3}  value={v!r:>15}  fmt={f!r}")


def main():
    p = argparse.ArgumentParser(
        description="Unveil — извлечение и расшифровка диапазона "
                    "из форматирования Excel-носителя CellVeilXL v3")
    p.add_argument("carrier",     help=".xlsx файл-носитель")
    p.add_argument("private_key", help="Base64URL-ключ, @path-к-файлу или $ENV_NAME")
    p.add_argument("output",      help="куда сохранить восстановленный .xlsx")
    p.add_argument("--mode", choices=["production", "demo"], default="production")
    p.add_argument("--sheet", default=None)
    p.add_argument("--start-cell", default="A1")
    p.add_argument("--print", dest="do_print", action="store_true")
    args = p.parse_args()

    pk = args.private_key
    if pk.startswith("@"):
        pk = Path(pk[1:]).read_text().strip()
    elif pk.startswith("$"):
        pk = os.environ[pk[1:]]
    mode_num = MODE_PROD if args.mode == "production" else MODE_DEMO

    print(f"[*] Открываю носитель: {args.carrier}")
    wb = load_workbook(args.carrier, rich_text=True)
    ws = wb[args.sheet] if args.sheet else wb.active

    print(f"[*] Извлекаю биты из форматирования (лист {ws.title!r})")
    reader = BitReader(ws)
    mode, payload = read_packet(reader)
    print(f"[*] Пакет: mode={mode}, encrypted_size={len(payload)} байт")

    if mode != mode_num:
        raise ValueError(f"Режим пакета ({mode}) ≠ выбранного ({mode_num}). "
                         "Проверьте --mode.")

    print("[*] Расшифровываю payload…")
    plain = cv_open(payload, pk, mode)
    print(f"[*] Расшифровано {len(plain)} байт")

    h, w, cells = deserialize(plain)
    print(f"[*] Восстановлено {h}×{w} = {h * w} ячеек")

    if args.do_print:
        print("[*] Содержимое:")
        print_cells(h, w, cells)

    print(f"[*] Сохраняю: {args.output}")
    write_output(h, w, cells, args.output, args.start_cell)
    print("[+] Готово")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"[!] Ошибка: {e}", file=sys.stderr)
        sys.exit(1)
