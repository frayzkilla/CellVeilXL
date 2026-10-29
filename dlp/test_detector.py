import json
import random
import subprocess
import sys
import tempfile
import unittest
import zipfile
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from dlp import Limits, analyze_xlsx
from dlp.fixtures import write_carrier
from dlp.demo import build_demo


class DetectorTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'document.xlsx'

    def test_packets(self):
        for mode in (1, 2):
            for magic in (b'CVX3', b'\x00\xff\x00\xff'):
                for shared in (False, True):
                    with self.subTest(mode=mode, magic=magic, shared=shared):
                        write_carrier(self.path, magic=magic, mode=mode, shared=shared,
                                      start=6, negative_rotation=True)
                        result = analyze_xlsx(self.path)
                        self.assertEqual(result['status'], 'ok', result)
                        self.assertGreaterEqual(result['risk_score'], 85, result)
                        packet = next(f for f in result['findings'] if 'magic_hex' in f)
                        self.assertEqual(packet['magic_hex'], magic.hex())
                        self.assertTrue(packet['complete'])

    def test_known_range_and_hidden_sheet(self):
        write_carrier(self.path, range_only=True, hidden=True, magic=b'ABCD')
        result = analyze_xlsx(self.path, host_range='B2:B16')
        self.assertEqual(result['verdict'], 'high_risk', result)
        self.assertEqual(result['sheets_scanned'], 2)

    def test_heuristic_without_header(self):
        write_carrier(self.path, signature=False)
        result = analyze_xlsx(self.path)
        self.assertEqual(result['verdict'], 'suspicious', result)
        self.assertIn('rgb_low_bits_pattern', [f['rule'] for f in result['findings']])
        self.assertEqual(result['profiles']['cellveilxl']['risk_score'], 0)

    def test_truncated_known_packet(self):
        write_carrier(self.path)
        wb = load_workbook(self.path, rich_text=True)
        wb.active.delete_rows(4, 20)
        wb.save(self.path)
        result = analyze_xlsx(self.path)
        packet = next(f for f in result['findings'] if f['rule'] == 'cellveilxl_packet')
        self.assertFalse(packet['complete'])
        self.assertEqual(packet['score'], 70)

    def test_visible_literal_is_not_a_packet(self):
        wb = Workbook()
        wb.active['A1'] = 'CVX3 is a string in a normal document'
        wb.save(self.path)
        self.assertEqual(analyze_xlsx(self.path)['risk_score'], 0)

    def test_ordinary_documents(self):
        for seed in range(20):
            wb = Workbook()
            ws = wb.active
            rng = random.Random(seed)
            for row in range(1, 40):
                cell = ws.cell(row, 1, 'Business report with regular formatting ' * 3)
                cell.font = Font(name=rng.choice(('Arial', 'Calibri', 'Tahoma', 'Courier New')),
                                 color=rng.choice(('FF203040', 'FF000000', 'FF0000FF')),
                                 bold=row % 5 == 0)
                cell.alignment = Alignment(textRotation=rng.choice((0, 0, 0, 15, 45)))
            ws['B1'] = '=SUM(1,2)'
            ws.merge_cells('C1:D3')
            wb.save(self.path)
            result = analyze_xlsx(self.path)
            self.assertEqual(result['risk_score'], 0, result)
            self.assertEqual(result['status'], 'ok')

    def test_limits(self):
        write_carrier(self.path)
        for field in ('file_bytes', 'expanded_bytes', 'zip_entries', 'cells', 'sheet_area', 'sheets', 'characters', 'findings'):
            with self.subTest(field=field):
                result = analyze_xlsx(self.path, limits=replace(Limits(), **{field: 0}))
                self.assertEqual(result['status'], 'limit_exceeded', result)
                self.assertEqual(result['verdict'], 'inconclusive')

    def test_invalid_and_missing_files(self):
        self.assertEqual(analyze_xlsx(self.path)['status'], 'error')
        self.path.write_bytes(b'not a zip archive')
        result = analyze_xlsx(self.path)
        self.assertEqual(result['verdict'], 'inconclusive')
        self.assertEqual(result['status'], 'error')

    def test_xml_entities_rejected(self):
        with zipfile.ZipFile(self.path, 'w') as archive:
            archive.writestr('[Content_Types].xml', '<Types/>')
            archive.writestr('xl/workbook.xml', '<!DOCTYPE a [<!ENTITY x "text">]><a>&x;</a>')
        result = analyze_xlsx(self.path)
        self.assertEqual(result['status'], 'error')
        self.assertEqual(result['verdict'], 'inconclusive')

    def test_demo_and_generic_profile_independence(self):
        output = Path(self.directory.name) / 'demo'
        records = build_demo(output)
        self.assertTrue(all(record['passed'] for record in records))
        unrelated = next(r for r in records if r['report'] == 'report-4.json')
        self.assertGreaterEqual(unrelated['generic_score'], 60)
        self.assertEqual(unrelated['cellveilxl_score'], 0)
        self.assertTrue((output / 'report.html').is_file())
        for record in records:
            report = json.loads((output / record['report']).read_text(encoding='utf-8'))
            self.assertGreaterEqual(report['duration_ms'], 0)

    def test_invalid_range(self):
        write_carrier(self.path)
        self.assertEqual(analyze_xlsx(self.path, host_range='bad range')['status'], 'error')

    def test_cli(self):
        write_carrier(self.path)
        result = subprocess.run([sys.executable, '-m', 'dlp', str(self.path)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 1, result.stderr)
        self.assertEqual(json.loads(result.stdout)['verdict'], 'high_risk')
        self.path.write_bytes(b'broken')
        result = subprocess.run([sys.executable, '-m', 'dlp', str(self.path)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(json.loads(result.stdout)['status'], 'error')
        wb = Workbook()
        wb.active['A1'] = 'Normal document'
        wb.save(self.path)
        result = subprocess.run([sys.executable, '-m', 'dlp', str(self.path)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0)


if __name__ == '__main__':
    unittest.main()
