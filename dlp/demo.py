import argparse
import html
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import TextBlock
from openpyxl.styles import Font

from .detector import analyze_xlsx
from .fixtures import write_carrier


def build_demo(output, include=()):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    normal = output / 'ordinary.xlsx'
    wb = Workbook()
    ws = wb.active
    for row in range(1, 80):
        ws.cell(row, 1, 'Quarterly business report and ordinary text')
        ws.cell(row, 2, row * 17)
        ws.cell(row, 1).font = Font(name='Calibri', bold=row % 10 == 0)
    wb.save(normal)
    cases = [('Обычный документ', normal, 'no_indicators')]
    for name, magic in [('standard', b'CVX3'), ('custom', b'\x00\xff\x01\xfe')]:
        path = output / (name + '.xlsx')
        write_carrier(path, magic=magic)
        cases.append(('Синтетический пакет CellVeilXL: ' + name, path, 'high_risk'))
    similar = output / 'unrelated_formatting.xlsx'
    write_carrier(similar, signature=False)
    wb = load_workbook(similar, rich_text=True)
    for row in wb.active:
        for cell in row:
            if cell.value is None:
                continue
            cell.font = Font(name='Courier New', color=cell.font.color)
            if isinstance(cell.value, list):
                for block in cell.value:
                    if isinstance(block, TextBlock):
                        block.font.rFont = 'Courier New'
    wb.save(similar)
    cases.append(('Похожий канал: другой шрифт, нет заголовка пакета', similar, 'suspicious'))
    broken = output / 'broken.xlsx'
    broken.write_bytes(b'Invalid XLSX for failure-policy demonstration')
    cases.append(('Нечитаемый входной файл', broken, 'inconclusive'))
    cases.extend(('Образец Excel/VBA: ' + Path(path).parent.name + '/' + Path(path).name, Path(path), None) for path in include)
    records = []
    rows = []
    for index, (label, path, expected) in enumerate(cases, 1):
        report = analyze_xlsx(path)
        report_name = f'report-{index}.json'
        (output / report_name).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        generic = report['profiles']['formatting_steganography']['risk_score']
        specific = report['profiles']['cellveilxl']['risk_score']
        rules = ', '.join(dict.fromkeys(f['rule'] for f in report['findings'])) or 'None'
        records.append({'case': label, 'expected': expected, 'actual': report['verdict'],
                        'passed': report['verdict'] == expected if expected else None,
                        'risk_score': report['risk_score'], 'generic_score': generic,
                        'cellveilxl_score': specific, 'status': report['status'], 'report': report_name})
        rows.append(f'<tr><td>{html.escape(label)}</td><td>{html.escape(report["status"])}</td>'
                    f'<td><b>{generic}</b></td><td><b>{specific}</b></td>'
                    f'<td>{html.escape(report["verdict"])}</td><td>{html.escape(rules)}</td>'
                    f'<td><a href="{report_name}">JSON</a></td></tr>')
    passed = sum(r['passed'] is True for r in records)
    checked = sum(r['passed'] is not None for r in records)
    document = '''<!doctype html><html lang="ru"><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'">
<title>Детекция скрытых данных в XLSX — демонстрация для DLP</title>
<style>body{font:16px/1.6 system-ui;background:#f4f6fa;color:#17263b;margin:0;padding:40px}
main{max-width:1350px;margin:auto}h1{font-size:36px;line-height:1.2}table{border-collapse:collapse;width:100%;background:white}
th,td{text-align:left;padding:14px;border-bottom:1px solid #dce3ee;vertical-align:top}th{background:#17263b;color:white}
.note{padding:20px;background:#e5ecf8;border-left:4px solid #355dad;margin:24px 0}a{color:#244d9a}
small{color:#4d6079}.scroll{overflow:auto}code{background:#e5ecf8;padding:3px 5px}</style><main>
<small>ВОСПРОИЗВОДИМАЯ ДЕМОНСТРАЦИЯ / БЕЗ КОНФИДЕНЦИАЛЬНЫХ ВХОДНЫХ ДАННЫХ</small>
<h1>Скрытые данные в оформлении Excel</h1>
<p>Независимый анализ объёма файла, дробления текста, младших битов RGB, наклонов и видимости.
Отдельный протокольный блок распознаёт потоки формата CellVeilXL.</p>
<div class="note">Баллы отражают правила риска, а не вероятность. Общая аномалия не устанавливает инструмент и не доказывает утечку.
Синтетические пакеты содержат тестовые байты. Образцы, созданные настоящим Excel/VBA, отмечены отдельно.</div>
<div class="scroll"><table><thead><tr><th>Документ</th><th>Статус проверки</th><th>Общий риск</th><th>Риск CellVeilXL</th>
<th>Вердикт</th><th>Сработавшие правила</th><th>Детали</th></tr></thead><tbody>'''
    document += ''.join(rows) + f'''</tbody></table></div><p>Ожидаемый результат получен в <b>{passed}/{checked}</b> синтетических сценариев.
Это проверка примеров, а не измерение точности на корпоративных документах.</p>
<h2>Как провести пилот с отделом защиты</h2><ol><li>Собрать согласованную размеченную выборку типовых рабочих книг и начать в режиме наблюдения.</li>
<li>Измерить ложные тревоги по типам документов и обнаружение вариантов носителей, включая повреждённые.</li>
<li>Настроить версионируемую политику и проверить её на отдельной выборке.</li>
<li>Включать карантин или ручную проверку через DLP. Ошибка сканирования должна обрабатываться отдельной политикой.</li></ol>
<p>Повторить демонстрацию: <code>python -m dlp.demo --output .test-output/dlp-demo</code><br>
Автотесты: <code>python -m unittest discover -s dlp -p "test_*.py" -v</code></p>
<p><small>Границы: рендеринг и расшифровка не выполняются; цвета тем и индексированные цвета не декодируются в RGB;
изменение раскладки ячеек может помешать восстановлению протокола. Действуют лимиты ZIP и XML;
промышленному рабочему процессу также нужны ограничения памяти и времени.</small></p></main></html>'''
    (output / 'report.html').write_text(document, encoding='utf-8')
    (output / 'summary.json').write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
    return records


def main():
    parser = argparse.ArgumentParser(description='Build reproducible XLSX DLP demonstration artifacts')
    parser.add_argument('--output', default='.test-output/dlp-demo')
    parser.add_argument('--include', nargs='*', default=[])
    args = parser.parse_args()
    records = build_demo(args.output, args.include)
    print(json.dumps(records, ensure_ascii=True, indent=2))
    return 1 if any(r['passed'] is False for r in records) else 0


if __name__ == '__main__':
    raise SystemExit(main())
