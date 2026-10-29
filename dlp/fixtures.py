import random
import zipfile

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font


def write_carrier(path, *, magic=b'CVX3', mode=2, start=1, range_only=False,
                  signature=True, shared=False, negative_rotation=False, hidden=False):
    rng = random.Random(1234)
    body = bytes(rng.randrange(1, 256) for _ in range(300))
    if mode == 1:
        body = pow(4, 12345, 154742504910672534362390399).to_bytes(11, 'little') + body[11:] + b'\x01'
    else:
        body = bytes.fromhex('8520f0098930a754748b7ddcb43ef75a0dbf3a0d26381af4eba4a98eaa9b4e6a') + body[32:]
    payload = magic + bytes((mode, 1)) + len(body).to_bytes(4, 'little') + body
    bits = ''.join(format(byte, '08b') for byte in payload)
    if not signature:
        bits = ''.join(str(rng.randrange(2)) for _ in bits)
    wb = Workbook()
    ws = wb.active
    if hidden:
        wb.create_sheet('Visible')
        ws.sheet_state = 'hidden'
    ws['A1'] = 'Ordinary bold heading'
    ws['A1'].font = Font(bold=True)
    offset = 0
    fonts = ('Arial', 'Calibri', 'Segoe UI', 'Tahoma')
    for row in range(start + 1, start + 16):
        if range_only:
            ws.cell(row, 1, 'Outside the host range').font = Font(name='Courier New')
        cell = ws.cell(row, 2, 'Normal report text with sufficient characters 0123456789')
        if offset >= len(bits):
            continue
        def take(n):
            nonlocal offset
            value = int(bits[offset:offset + n].ljust(n, '0'), 2)
            offset += n
            return value
        rotation = take(2)
        rotation = 90 - (-8 + rotation) if negative_rotation else rotation
        cell.alignment = Alignment(textRotation=rotation)
        name = fonts[take(2)]
        cell.font = Font(name=name, color='FF3C5068')
        rich = []
        for char in str(cell.value):
            if offset >= len(bits):
                rich.append(TextBlock(InlineFont(rFont=name, color=None), char))
                cell.font = Font(name=name, color=None)
                continue
            value = take(6)
            red, green, blue = 60 + (value >> 4), 80 + ((value >> 2) & 3), 104 + (value & 3)
            rich.append(TextBlock(InlineFont(rFont=name, color=f'FF{red:02X}{green:02X}{blue:02X}'), char))
        cell.value = CellRichText(rich)
    wb.save(path)
    if shared:
        from xml.etree import ElementTree as ET
        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        with zipfile.ZipFile(path) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        sheet = ET.fromstring(files['xl/worksheets/sheet1.xml'])
        strings = ET.Element(ns + 'sst')
        for cell in sheet.iter(ns + 'c'):
            inline = cell.find(ns + 'is')
            if inline is None:
                continue
            cell.remove(inline)
            item = ET.SubElement(strings, ns + 'si')
            item.extend(list(inline))
            cell.set('t', 's')
            ET.SubElement(cell, ns + 'v').text = str(len(strings) - 1)
        files['xl/worksheets/sheet1.xml'] = ET.tostring(sheet)
        files['xl/sharedStrings.xml'] = ET.tostring(strings)
        types = ET.fromstring(files['[Content_Types].xml'])
        ET.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                      PartName='/xl/sharedStrings.xml', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
        files['[Content_Types].xml'] = ET.tostring(types)
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, data in files.items():
                archive.writestr(name, data)


